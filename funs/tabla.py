from pptx.table import Table, _Row, _Column, _Cell
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from copy import deepcopy

def add_row(table: Table) -> _Row:
    """
    Crea una fila nueva
    No sé como funciona, pero funciona
    """
    new_row = deepcopy(table._tbl.tr_lst[-1]) 
    # duplicating last row of the table as a new row to be added

    for tc in new_row.tc_lst:
        cell = _Cell(tc, new_row.tc_lst)
        cell.text = '' # defaulting cell contents to empty text

    table._tbl.append(new_row) 
    return table.rows[len(table.rows)-1]


def cell_set(cell, cell_text=str, text_font='Calibri', text_size=14, alignment_center=True, text_color=RGBColor(0, 0, 0), bold=False):
    """
    Define el estilo del texto de las celdas
    """
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if alignment_center == True:
        cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = cell.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = cell_text
    font = run.font
    font.name = text_font
    font.size = Pt(text_size)
    font.color.rgb = text_color
    if bold == True:
        font.bold = True

def heath(cell, glo, raw):
    """
    Rellena de colores las celdas que tengan 10 o más puntos de diferencia con el promedio global
    Redondea el valor glob para que la diferencia en los valores mencionados en las tablas sean congruentes con las diferencias

    cell = Celda elegida
    glo = Promedio global de la dimensión específica
    raw = Puntaje de la celda que se va a comparar
    """
    
    glo = round(glo*100)
    raw = round(raw*100)
    
    dif = raw - glo
    
    # Diferencia Positiva
    if (dif >= 10) and (dif < 15):
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(215, 228, 189)
    elif dif >= 15:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(173, 219, 123)

    # Diferencia Negativa
    elif (dif <= -10) and (dif > -15):
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(253, 212, 181)
    elif dif <= -15:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(242, 160, 104)